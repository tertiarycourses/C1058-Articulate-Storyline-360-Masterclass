import re, sys
from pptx import Presentation
prs = Presentation(sys.argv[1])
leak = re.compile(r"TGS-?|(?<!NON-)WSQ|SSG|SkillsFuture|TRAQOM|assess|attendance|\bfund(ing|ed)?\b|subsid|\bLO[0-9]|\b[KA][1-9]\b|open.book|exam|competen|TSC|Skills Framework|6:30|18:30|four.day|4 days|32 hours|candidate", re.I)
for n, s in enumerate(prs.slides):
    for sh in s.shapes:
        texts = []
        if sh.has_text_frame: texts.append(sh.text_frame.text)
        if sh.has_table: texts += [c.text for r in sh.table.rows for c in r.cells]
        for t in texts:
            if 'Mastering Articulate Storyline 360 for e-Learning Content Creation  ·  TGS' in t or 'Articulate Storyline 360 Masterclass  ·  C1058' in t: continue
            for m in leak.finditer(t):
                print(f"[{n}] {sh.name}: {t[max(0,m.start()-70):m.end()+50]!r}"); break
