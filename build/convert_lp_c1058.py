"""C1058 Lesson Plan: WSQ 4-day TGS-2024042308 LP -> non-WSQ 2-day C1058 LP (edited in place)."""
import copy, json
from docx import Document

SRC = "source-wsq/courseware/LP-Mastering Articulate Storyline 360 for e-Learning Content Creation.docx"
OUT = "courseware/LP-Articulate Storyline 360 Masterclass.docx"
OLD_T, NEW_T = "Mastering Articulate Storyline 360 for e-Learning Content Creation", "Articulate Storyline 360 Masterclass"
smap = {int(k): v for k, v in json.load(open("build/slide_map.json")).items()}


def sl(*olds):
    return ", ".join(str(smap[o]) for o in olds)


def rng(a, b):
    kept = [smap[n] for n in range(a, b + 1) if n in smap]
    return f"Slides {min(kept)}–{max(kept)}"


def set_par(p, text):
    if not p.runs:
        p.add_run(text)
        return
    p.runs[0].text = text
    for r in p.runs[1:]:
        r.text = ""


def set_cell(c, text):
    set_par(c.paragraphs[0], text)


def drop(el):
    el.getparent().remove(el)


d = Document(SRC)
T = d.tables
proto = copy.deepcopy(T[2].rows[1]._tr)          # body rows are unshaded; any body row is the prototype

T1 = rng(17, 45)
T2 = rng(46, 78)
T3 = rng(79, 107)
T4 = rng(108, 133)
DAY1 = [  # (start, minutes, mode, activity, reference) — 9:30-17:30, 450 teaching min + 30-min lunch
    ("09:30", 15, "Admin", "Welcome, introductions, course materials and course outline", rng(1, 16)),
    ("09:45", 90, "Topic 1", "Performance gaps, digital storytelling, content promise and storyboard architecture; project model, interface, Story View, scenes, slides and layers", f"Topic 1 begins at slide {smap[17]}"),
    ("11:15", 45, "Topic 1", "Prototype reviews, question slides, result logic, question banks and media libraries", T1),
    ("12:00", 30, "Labs 1–2", "Theme/media replacement and storyboard-content build", f"Lab slides {sl(43, 44)}"),
    ("12:30", 30, "Lunch", "Lunch break", "—"),
    ("13:00", 45, "Labs 1–2", "Labs 1–2 continued; guided Storyline practice and Topic 1 checkpoint", f"Topic 1 checkpoint · slide {smap[45]}"),
    ("13:45", 90, "Topic 2", "Audience preferences, visual grammar, pictures, 360° media and video; video workflow, audio, captions, accessible equivalents, web objects and screen recordings", f"Topic 2 begins at slide {smap[46]}"),
    ("15:15", 30, "Topic 2", "Native objects, text, tables, formatting and layout QA", T2),
    ("15:45", 75, "Labs 3–4", "Audience-fit media curation and governed AI-assisted variation", f"Lab slides {sl(74, 75)}"),
    ("17:00", 30, "Checkpoint", "Guided media/accessibility practice, provenance review and Day 1 recap", "Topic 2 checkpoint"),
]
DAY2 = [
    ("09:30", 15, "Admin", "Day 1 retrieval practice and Day 2 outline", "Recap"),
    ("09:45", 45, "Labs 5–6", "Narration, captions and scene-wide accessibility; Topic 2 checkpoint", f"Lab slides {sl(76, 77)} · checkpoint slide {smap[78]}"),
    ("10:30", 60, "Topic 3", "Meaningful choices, learner-controlled discovery, triggers, states and variables", f"Topic 3 begins at slide {smap[79]}"),
    ("11:30", 60, "Topic 3", "Timeline, cue points, animation, freeform questions and branching", T3),
    ("12:30", 30, "Lunch", "Lunch break", "—"),
    ("13:00", 90, "Labs 7–8", "Graded question, multi-question quiz and result logic; guided interaction testing and Topic 3 checkpoint", f"Lab slides {sl(105, 106)} · checkpoint slide {smap[107]}"),
    ("14:30", 60, "Topic 4", "Themes, masters, player, multi-device UX, accessibility and Review 360", f"Topic 4 begins at slide {smap[108]}"),
    ("15:30", 60, "Topic 4", "Preview strategy, publishing channels, LMS/LRS standards and tracking; distribution frequency, execution guidelines, lifecycle and analytics", T4),
    ("16:30", 45, "Lab 9", "Summary, Review 360 and controlled publish plan; Topic 4 checkpoint", f"Lab slide {smap[132]} · checkpoint slide {smap[133]}"),
    ("17:15", 15, "Close", "Course synthesis, Q&A and course feedback", rng(134, 142)),
]


def hhmm(t, add):
    h, m = map(int, t.split(":"))
    m += add
    return f"{h + m // 60:02d}:{m % 60:02d}"


def rebuild(table, rows):
    tbl = table._tbl
    for tr in [r._tr for r in table.rows][1:]:
        tbl.remove(tr)
    teach, end = 0, "09:30"
    for start, mins, mode, text, ref in rows:
        assert start == end, (start, end, text)
        tbl.append(copy.deepcopy(proto))
        c = table.rows[-1].cells
        end = hhmm(start, mins)
        for cell, v in zip(c, [f"{start}–{end}", f"{mins} min", mode, text, ref]):
            set_cell(cell, v)
        if mode != "Lunch":
            teach += mins
    assert teach == 450 and end == "17:30", (teach, end)


rebuild(T[2], DAY1)
rebuild(T[3], DAY2)
drop(T[4]._tbl)
drop(T[5]._tbl)

# ---- cover ----
for p in d.paragraphs:
    t = p.text.strip()
    if t == OLD_T:
        set_par(p, NEW_T)
    elif t == "TGS Ref No: TGS-2024042308":
        set_par(p, "Course Code: C1058")
    elif t == "Version 10":
        set_par(p, "Version 1.0")

# ---- version record: a single initial-release row ----
vt = T[0]
r = vt.rows[1]
for cell, v in zip(r.cells, ["1.0", "2 October 2026",
                             "Initial release — two-day plan covering four topics, 48 Storyline concepts and nine .story labs.",
                             "Dr Alfred Ang"]):
    set_cell(cell, v)
for row in list(vt.rows)[2:]:
    drop(row._tr)

# ---- course information ----
for row in T[1].rows:
    k = row.cells[0].text.strip()
    if k == "Course title":
        set_cell(row.cells[1], NEW_T)
    elif k == "Course code":
        set_cell(row.cells[1], "C1058")
    elif k == "Duration":
        set_cell(row.cells[1], "2 days · 15 hours (7.5 instructional hours per day)")
    elif k == "Delivery":
        set_cell(row.cells[1], "9:30am – 5:30pm daily with a 30-minute lunch · concept facilitation + hands-on Storyline labs")
    elif k == "Version":
        set_cell(row.cells[1], "v1.0 · 2 October 2026")
for row in list(T[1].rows):
    if row.cells[0].text.strip() == "TSC":
        drop(row._tr)

# ---- lab alignment: keep topic + LO, drop the K/A competency codes, remap slides ----
for row in list(T[6].rows)[1:]:
    topic = row.cells[1].text.split(" · ")
    set_cell(row.cells[1], " · ".join(topic[:2]))
    set_cell(row.cells[2], str(smap[int(row.cells[2].text.strip())]))
    if "assess the mapped criteria" in row.cells[3].text:
        set_cell(row.cells[3], row.cells[3].text.replace("assess the mapped criteria", "test the taught decisions"))

# ---- body: drop Assessment section + Day 3/4, retitle schedule ----
paras = d.paragraphs
in_assess = False
for p in list(paras):
    t = p.text.strip()
    st = p.style.name
    if st == "Heading 1" and t == "Assessment":
        in_assess = True
        drop(p._p)
        continue
    if in_assess:
        if st.startswith("Heading"):
            in_assess = False
        else:
            drop(p._p)
            continue
    if (st == "Heading 2" and t in ("Day 3", "Day 4")) or t.startswith(("Assessment\t", "Day 3\t", "Day 4\t")):
        drop(p._p)
    elif t == "Four-Day Schedule":
        set_par(p, "Two-Day Schedule")
    elif t.startswith("Four-Day Schedule\t"):
        set_par(p, p.text.replace("Four-Day", "Two-Day"))
    elif st == "Heading 2" and t in ("Day 1", "Day 2"):
        set_par(p, f"{t} · 9:30am – 5:30pm")

# ---- TOC page numbers, re-read from the rendered PDF (Word's TOC is static text here) ----
TOC = {"Course Information": 4, "Learning Outcomes": 4, "Two-Day Schedule": 4, "Day 1": 4, "Day 2": 5, "Lab Alignment": 7}
for p in d.paragraphs:
    if "\t" in p.text and p.text.split("\t")[1].strip().isdigit():
        for k, pg in TOC.items():
            if p.text.startswith(k):
                set_par(p, p.text.split("\t")[0] + "\t" + str(pg))

d.save(OUT)
for p in d.paragraphs:
    if p.text.strip():
        print(p.style.name, "|", p.text[:120])
for t in d.tables:
    print("----")
    for r in t.rows:
        print(" | ".join(c.text for c in r.cells)[:170])
