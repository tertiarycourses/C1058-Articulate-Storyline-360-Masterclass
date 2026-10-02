"""C1058 Learner Guide: WSQ TGS-2024042308 LG -> non-WSQ C1058 LG (edited in place)."""
from docx import Document

SRC = "source-wsq/courseware/LG-Mastering Articulate Storyline 360 for e-Learning Content Creation.docx"
OUT = "courseware/LG-Articulate Storyline 360 Masterclass.docx"
OLD_T, NEW_T = "Mastering Articulate Storyline 360 for e-Learning Content Creation", "Articulate Storyline 360 Masterclass"
URL = "https://www.tertiarycourses.com.sg/articulate-storyline-360-essential-training.html"


def set_par(p, text):
    p.runs[0].text = text
    for r in p.runs[1:]:
        r.text = ""


def set_cell(c, text):
    set_par(c.paragraphs[0], text)


def drop(el):
    el.getparent().remove(el)


d = Document(SRC)
T = d.tables

# ---- cover ----
for p in d.paragraphs:
    t = p.text.strip()
    if t == OLD_T:
        set_par(p, NEW_T)
    elif t == "TGS Ref No: TGS-2024042308":
        set_par(p, "Course Code: C1058")
    elif t == "Version 10":
        set_par(p, "Version 1.0")

# ---- version record: a single initial-release row ----
vt = T[0]
for cell, v in zip(vt.rows[1].cells, ["1.0", "2 October 2026",
                                      "Initial release — guide covering 48 Storyline concepts and nine self-contained .story labs with detailed procedures, evidence and acceptance criteria.",
                                      "Dr Alfred Ang"]):
    set_cell(cell, v)
for row in list(vt.rows)[2:]:
    drop(row._tr)

# ---- source register: non-WSQ course page; drop WSQ deck + assessment rows ----
sr = T[2]
for row in list(sr.rows)[1:]:
    k = row.cells[0].text.strip()
    if k == "Official course page":
        set_cell(row.cells[1], URL)
    elif k.startswith(("Legacy trainer deck", "Drive assessment", "Approved assessment", "Approved courseware mapping")):
        drop(row._tr)

# ---- TSC Alignment section (heading + K/A table) ----
drop(T[1]._tbl)
for p in list(d.paragraphs):
    t = p.text.strip()
    if t == "TSC Alignment" or t.startswith("TSC Alignment\t"):
        drop(p._p)
    elif t.startswith("Alignment: LO"):
        set_par(p, "Learning outcome: " + t.split("Alignment: ")[1].split(" · ")[0])
    elif "and formal assessment" in t:
        set_par(p, p.text.replace("and formal assessment", "and graded quiz release"))
    elif "assess the mapped criteria" in t:
        set_par(p, p.text.replace("assess the mapped criteria", "test the taught decisions"))

# ---- TOC page numbers, re-read from the rendered PDF (Word's TOC is static text here) ----
TOC = {"Purpose and Use": 4, "Learning Outcomes": 4, "Core Professional Workflow": 4, "Topic 1": 5, "Topic 2": 11, "Topic 3": 21, "Topic 4": 28, "Intellectual Property": 33, "Publishing and Lifecycle": 33, "Source Register": 33}
for p in d.paragraphs:
    if "\t" in p.text and p.text.split("\t")[1].strip().isdigit():
        for k, pg in TOC.items():
            if p.text.startswith(k):
                set_par(p, p.text.split("\t")[0] + "\t" + str(pg))

d.save(OUT)
for p in d.paragraphs[:40]:
    if p.text.strip():
        print(p.style.name, "|", p.text[:110])
