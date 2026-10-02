"""C1058 deck: WSQ TGS-2024042308 (4 days) -> non-WSQ C1058 (2 days), edited in place (run-level, formatting kept)."""
import json, re
from pptx import Presentation

SRC = "source-wsq/courseware/Mastering Articulate Storyline 360 for e-Learning Content Creation-v10.pptx"
OUT = "courseware/Articulate Storyline 360 Masterclass-v1.0.pptx"
OLD_T, NEW_T = "Mastering Articulate Storyline 360 for e-Learning Content Creation", "Articulate Storyline 360 Masterclass"
OLD_C, NEW_C = "TGS-2024042308", "C1058"
# 0-based source indexes: digital attendance, Skills Framework alignment, K criteria, A criteria,
# Learning Outcomes, Assessment Design, Assessment Briefing, Final Assessment, Assessment Flow
DROP = [1, 7, 8, 9, 10, 12, 13, 138, 139]
FEEDBACK = 140                               # closing digital-attendance slide -> Course Feedback


def para_replace(p, old, new):
    if old not in p.text:
        return
    for r in p.runs:
        if old in r.text:
            r.text = r.text.replace(old, new)
    if old in p.text:                       # spans runs: collapse into the first run
        full = p.text.replace(old, new)
        p.runs[0].text = full
        for r in p.runs[1:]:
            r.text = ""


def set_paras(shape, lines):
    """Write line i into paragraph i's first run, keeping each paragraph's own format."""
    paras = list(shape.text_frame.paragraphs)
    assert len(paras) >= len(lines), (shape.name, len(paras), lines)
    for p, line in zip(paras, lines):
        p.runs[0].text = line
        for r in p.runs[1:]:
            r.text = ""
    for p in paras[len(lines):]:
        p._p.getparent().remove(p._p)


def shape_by_text(slide, startswith):
    hits = [s for s in slide.shapes if s.has_text_frame and s.text_frame.text.strip().startswith(startswith)]
    assert len(hits) == 1, (startswith, len(hits))
    return hits[0]


def run_replace(slide, old, new):
    """Replace inside a single run (multi-line card text kept as one run with literal newlines)."""
    n = 0
    for s in slide.shapes:
        if s.has_text_frame:
            for p in s.text_frame.paragraphs:
                for r in p.runs:
                    if old in r.text:
                        r.text = r.text.replace(old, new)
                        n += 1
    assert n == 1, (old, n)


prs = Presentation(SRC)
S = prs.slides

# --- cover ---
set_paras(shape_by_text(S[0], "COURSE SLIDES"), ["COURSE SLIDES · NON-WSQ"])
set_paras(shape_by_text(S[0], "Course code"),
          ["Course code  C1058", "Provider  Tertiary Infotech Academy Pte Ltd · UEN: 201200696W"])
set_paras(shape_by_text(S[0], "Version v10"), ["Version v1.0 · 2 October 2026"])

# --- ground rules / readiness / LMS package / course flow ---
set_paras(shape_by_text(S[4], "Protect learner data"), ["Protect learner data and course materials."])
set_paras(shape_by_text(S[14], "Versioned .story files"), ["Versioned .story files, protected source content and verified publishes."])
set_paras(shape_by_text(S[5], "Courseware · v10"), ["Courseware · v1.0"])
set_paras([s for s in S[5].shapes if s.has_text_frame and s.text_frame.text.strip() == "Mastering Articulate Storyline 360"][0], [NEW_T])
set_paras(shape_by_text(S[6], "Four-Day Learning Journey"), ["Two-Day Learning Journey"])

# --- course outline cards + topic dividers: competency codes -> delivery day ---
DAYS = {"LO1 · A1 · A2 · K1 · K3": "Day 1", "LO2 · A3 · A5 · K2 · K5": "Day 1 · Labs 5–6 on Day 2",
        "LO3 · A6 · K6": "Day 2", "LO4 · A4 · A7 · K4": "Day 2"}
for old, new in DAYS.items():
    run_replace(S[11], "\n" + old, "\n" + new)
for idx, old in [(16, "LO1 · A1 · A2 · K1 · K3"), (45, "LO2 · A3 · A5 · K2 · K5"),
                 (78, "LO3 · A6 · K6"), (107, "LO4 · A4 · A7 · K4")]:
    set_paras(shape_by_text(S[idx], old), [DAYS[old]])

# --- Interaction QA card: e-learning wording that reads as a WSQ "formal assessment" ---
set_paras(shape_by_text(S[102], "Before stakeholder review"), ["Before stakeholder review, LMS upload and graded quiz release."])

# --- execution-control eyebrow + support slide ---
set_paras(shape_by_text(S[136], "A7 · EXECUTION CONTROL"), ["EXECUTION CONTROL"])
set_paras(shape_by_text(S[137], "Do not send assessment"), ["Do not share learner data or course files publicly."])

# --- closing digital-attendance slide -> Course Feedback (rewritten in place) ---
fb = S[FEEDBACK]
set_paras(shape_by_text(fb, "CLOSING ADMIN"), ["BEFORE YOU LEAVE"])
set_paras(shape_by_text(fb, "Digital Attendance"), ["Course Feedback"])
for start, line in [("Assessment event", "Share your feedback"),
                    ("Complete the official", "Please complete the course feedback form before you leave."),
                    ("Verification", "Certificate"),
                    ("Remain until", "Your certificate of completion is emailed after the course."),
                    ("Feedback", "Keep in touch"),
                    ("Use the official", "Questions after the course? Email enquiry@tertiaryinfotech.com.")]:
    set_paras(shape_by_text(fb, start), [line])

# --- title / code everywhere ---
for slide in S:
    for shp in slide.shapes:
        frames = []
        if shp.has_text_frame:
            frames.append(shp.text_frame)
        if shp.has_table:
            frames += [c.text_frame for r in shp.table.rows for c in r.cells]
        for tf in frames:
            for p in tf.paragraphs:
                para_replace(p, OLD_T, NEW_T)
                para_replace(p, OLD_C, NEW_C)

# --- drop WSQ slides (highest first) ---
keep = [i for i in range(len(S)) if i not in DROP]
lst = prs.slides._sldIdLst
items = list(lst)
for idx in sorted(DROP, reverse=True):
    rid = items[idx].get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id")
    prs.part.drop_rel(rid)
    lst.remove(items[idx])

# --- renumber static bottom-right footers ---
H, W = prs.slide_height, prs.slide_width
for n, slide in enumerate(prs.slides, start=1):
    c = [s for s in slide.shapes if s.has_text_frame and s.text_frame.text.strip().isdigit()
         and s.top is not None and s.top > H * 0.85 and s.left > W * 0.5]
    if c:
        set_paras(max(c, key=lambda s: (s.left, s.top)), [str(n)])

prs.save(OUT)
# old 1-based slide number -> new 1-based slide number (for LP/LG slide references)
json.dump({str(o + 1): k + 1 for k, o in enumerate(keep)}, open("build/slide_map.json", "w"))

leak = re.compile(r"TGS|(?<!NON-)WSQ|SSG|SkillsFuture|TRAQOM|assess|attendance|\bfund(ing|ed)?\b|subsid|\bLO[0-9]|\b[KA][1-9]\b|open.book|competen|TSC|four.day|32 hours", re.I)
for n, slide in enumerate(prs.slides, start=1):
    txt = " // ".join(s.text_frame.text for s in slide.shapes if s.has_text_frame)
    for m in leak.finditer(txt):
        print(f"slide {n}: ...{txt[max(0, m.start()-60):m.end()+40]!r}")
print("wrote", OUT, len(prs.slides), "slides")
