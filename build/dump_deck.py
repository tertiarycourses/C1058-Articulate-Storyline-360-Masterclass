import sys
from pptx import Presentation
from pptx.util import Emu
prs = Presentation(sys.argv[1]); H, W = prs.slide_height, prs.slide_width
want = None if len(sys.argv) < 3 else {int(x) for x in sys.argv[2].split(',')}
for i, s in enumerate(prs.slides):
    if want is not None and i not in want: continue
    print(f"===== [{i}]")
    for sh in s.shapes:
        if sh.has_text_frame and sh.text_frame.text.strip():
            ps = [p.text for p in sh.text_frame.paragraphs]
            print(f"  T {sh.shape_id} {sh.name!r} np={len(ps)} @({sh.left*100//W},{sh.top*100//H}): {ps}")
        elif sh.has_table:
            for r in sh.table.rows: print("  TBL |", " | ".join(c.text for c in r.cells))
        elif sh.shape_type == 13:
            print(f"  PIC {sh.name} {sh.width*100//W}x{sh.height*100//H}% @({sh.left*100//W},{sh.top*100//H})")
